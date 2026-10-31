"""Excel (.xlsx) sheets with openpyxl, rows tab-separated.

Locations are "sheet 'Builds' rows 201–400". The cell cap
(``KNOWLEDGE_BASE_SHEET_MAX_CELLS``) is shared by every sheet of the file and
by the .xls reader. Cached formula results are read (``data_only``), never
the formulas. A cell's text is cut at ``CELL_MAX_CHARS`` and a row stops
once it would pass the item's remaining character budget, so one long shared
string repeated across a row cannot grow past it.
"""

from __future__ import annotations

import io
import zipfile
from collections.abc import Iterable, Iterator
from datetime import date, datetime, time, timedelta

from django.conf import settings
from openpyxl import load_workbook
from openpyxl.worksheet._read_only import ReadOnlyWorksheet

from knowledge_base.archive_guard import check_ooxml_container
from knowledge_base.chunking import (
    ELLIPSIS,
    SectionCollector,
    format_number,
    quoted,
    range_label,
)
from knowledge_base.exceptions import UnsupportedContent
from knowledge_base.text_limits import CHUNK_TARGET_CHARS

DAMAGED_MESSAGE = "This Excel workbook is damaged or could not be read."
# Longest text kept from one cell (an ellipsis marks the cut).
CELL_MAX_CHARS = 4000
# A sheet row as (row number, cell count, lazily formatted cell texts).
SheetRow = tuple[int, int, Iterator[str]]


class CellBudget:
    """Counts cells read across all sheets of one file."""

    def __init__(self) -> None:
        self.max_cells: int = settings.KNOWLEDGE_BASE_SHEET_MAX_CELLS
        self.used = 0

    def take(self, count: int) -> bool:
        """Count ``count`` cells; False once the cap is passed."""
        self.used += max(count, 1)
        return self.used <= self.max_cells

    def warning(self) -> str:
        return f"Only the first {format_number(self.max_cells)} cells were read."


def format_cell(value: object) -> str:
    """One cell as text: dates in ISO form, whole floats without ".0"."""
    if value is None:
        return ""
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, float):
        return str(int(value)) if value.is_integer() else repr(value)
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, (date, time)):
        return value.isoformat()
    if isinstance(value, timedelta):
        return str(value)
    text = str(value)
    if len(text) <= CELL_MAX_CHARS:
        return " ".join(text.split())
    # Only a prefix is normalized, never the whole (possibly huge) string.
    return " ".join(text[:CELL_MAX_CHARS].split())[: CELL_MAX_CHARS - 1] + ELLIPSIS


def add_sheet_rows(
    sheet_name: str,
    rows: Iterable[SheetRow],
    cells: CellBudget,
    collector: SectionCollector,
) -> bool:
    """Add ``(row number, cell count, cell texts)`` rows as sections of about
    a chunk.

    Returns False when the collector is full or the cell cap was reached.
    """
    prefix = f"sheet {quoted(sheet_name, 40)} "
    buffer: list[tuple[int, str]] = []
    size = 0
    for number, count, texts in rows:
        if not cells.take(count):
            collector.warn(cells.warning())
            _flush(prefix, sheet_name, buffer, collector)
            return False
        values, complete = _row_values(texts, collector.room - size)
        if not complete:
            # The row alone passes the item budget: keep what fits and stop.
            collector.truncated = True
            buffer.append((number, "\t".join(values)))
            _flush(prefix, sheet_name, buffer, collector)
            return False
        while values and not values[-1]:
            values.pop()
        if not values:
            continue
        line = "\t".join(values)
        if buffer and size + len(line) + 1 > CHUNK_TARGET_CHARS:
            if not _flush(prefix, sheet_name, buffer, collector):
                return False
            buffer = []
            size = 0
        buffer.append((number, line))
        size += len(line) + 1
    return _flush(prefix, sheet_name, buffer, collector)


def _row_values(texts: Iterator[str], room: int) -> tuple[list[str], bool]:
    """A row's cell texts while their tab-joined length fits ``room``.

    The flag is False when the row was cut there.
    """
    values: list[str] = []
    length = -1
    for text in texts:
        length += len(text) + 1
        if length > room:
            values.append(text[: max(room - (length - len(text)), 0)])
            return values, False
        values.append(text)
    return values, True


def _flush(
    prefix: str,
    sheet_name: str,
    buffer: list[tuple[int, str]],
    collector: SectionCollector,
) -> bool:
    if not buffer:
        return not collector.full
    location = prefix + range_label("row", "rows", buffer[0][0], buffer[-1][0])
    return collector.add(location, sheet_name, "\n".join(line for _, line in buffer))


def _rows(sheet: ReadOnlyWorksheet) -> Iterable[SheetRow]:
    # Missing rows are yielded empty by openpyxl, so numbering stays exact.
    for number, row in enumerate(sheet.values, start=1):
        yield number, len(row), (format_cell(value) for value in row)


def extract_xlsx(data: bytes, collector: SectionCollector) -> None:
    check_ooxml_container(data)
    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (zipfile.BadZipFile, KeyError, ValueError, TypeError, SyntaxError) as exc:
        # defusedxml / lxml parse errors are SyntaxError subclasses.
        raise UnsupportedContent(DAMAGED_MESSAGE) from exc
    cells = CellBudget()
    try:
        for name in workbook.sheetnames:
            sheet = workbook[name]
            if not isinstance(sheet, ReadOnlyWorksheet):
                continue
            # A wrong declared dimension would pad every row to its width.
            sheet.reset_dimensions()
            if not add_sheet_rows(name, _rows(sheet), cells, collector):
                return
    except (KeyError, ValueError, TypeError, SyntaxError) as exc:
        raise UnsupportedContent(DAMAGED_MESSAGE) from exc
    finally:
        workbook.close()
